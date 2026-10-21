"""Cadastro do modelo CEMIG — TCA ITMU a partir do pacote de referência.

O pacote traz a planilha original com os dados de um ensaio real. O cadastro
nunca usa esse arquivo direto: a cópia registrada é gerada aqui, esvaziando
somente os destinos variáveis e preservando rótulos, fórmulas, estilos,
mesclagens, dimensões, área de impressão e as tabelas auxiliares de correção.

Campos com significado ainda em dúvida são cadastrados para revisão humana e
não recebem mapeamento no Excel; sem mapeamento, a exportação não escreve nada
naquelas células.
"""

import json
from io import BytesIO
from pathlib import Path

from django.core.files.base import ContentFile
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from exportacao.services import build_cleaned_workbook

from . import tca_itmu_catalog as catalog
from .models import DocumentTemplate, ExcelFieldMapping, ExcelTemplate, TemplateField

CLEAN_WORKBOOK_NAME = "modelo-tca-itmu-limpo.xlsx"
EXCEL_TEMPLATE_NAME = "TCA ITMU — modelo limpo"


class PackageError(ValueError):
    """Pacote de referência incompleto ou incompatível com o catálogo."""


def reference_bytes(package_dir):
    path = Path(package_dir) / catalog.WORKBOOK_FILENAME
    if not path.is_file():
        raise PackageError("Arquivo {} não encontrado em {}.".format(catalog.WORKBOOK_FILENAME, package_dir))
    return path.read_bytes()


def _cells_in(sheet, reference):
    min_col, min_row, max_col, max_row = range_boundaries(reference)
    return [
        sheet.cell(row, col)
        for row in range(min_row, max_row + 1)
        for col in range(min_col, max_col + 1)
    ]


def _is_input(cell, preserved):
    """Entrada variável: desbloqueada, sem fórmula, fora dos textos fixos."""
    if isinstance(cell, MergedCell) or cell.protection.locked:
        return False
    if cell.data_type == "f":
        return False
    return cell.coordinate not in preserved


def preserved_coordinates(sheet):
    coordinates = {}
    for item in catalog.PRESERVED_TEXT:
        for cell in _cells_in(sheet, item["range"]):
            coordinates[cell.coordinate] = item["reason"]
    return coordinates


def audit(package_dir):
    """Confere o catálogo contra a planilha e devolve o relatório da revisão."""
    raw = reference_bytes(package_dir)
    workbook = load_workbook(BytesIO(raw), data_only=False)
    if catalog.SHEET_NAME not in workbook.sheetnames:
        raise PackageError("A planilha de referência não tem a aba {}.".format(catalog.SHEET_NAME))
    sheet = workbook[catalog.SHEET_NAME]
    preserved = preserved_coordinates(sheet)
    print_area = range_boundaries(catalog.PRINT_AREA)

    problems = []
    for item in catalog.CONFIRMED + catalog.PENDING:
        cell = sheet[item["cell"]]
        if isinstance(cell, MergedCell):
            problems.append("{}: {} é célula mesclada secundária.".format(item["field_id"], item["cell"]))
            continue
        if cell.protection.locked:
            problems.append("{}: {} está protegida no modelo.".format(item["field_id"], item["cell"]))
        if cell.data_type == "f":
            problems.append("{}: {} contém fórmula.".format(item["field_id"], item["cell"]))
        if cell.coordinate in preserved:
            problems.append("{}: {} guarda texto fixo da ficha.".format(item["field_id"], item["cell"]))
        min_col, min_row, max_col, max_row = print_area
        if not (min_row <= cell.row <= max_row and min_col <= cell.column <= max_col):
            problems.append("{}: {} está fora da área de impressão.".format(item["field_id"], item["cell"]))

    for item in catalog.PRESERVED_TEXT:
        if item["text"] is None:
            continue
        for cell in _cells_in(sheet, item["range"]):
            if cell.value != item["text"]:
                problems.append("Texto fixo {} mudou: {!r} no lugar de {!r}.".format(cell.coordinate, cell.value, item["text"]))

    inputs = [cell.coordinate for cell in _cells_in(sheet, catalog.PRINT_AREA) if _is_input(cell, preserved)]
    mapped = {item["cell"] for item in catalog.CONFIRMED}
    pending = {item["cell"] for item in catalog.PENDING}
    candidates = _candidate_cells(package_dir)

    return {
        "problems": problems,
        "sheet": catalog.SHEET_NAME,
        "formulas": sum(1 for row in sheet.iter_rows() for cell in row if cell.data_type == "f"),
        "merges": len(sheet.merged_cells.ranges),
        "print_area": str(sheet.print_area),
        "confirmed": len(catalog.CONFIRMED),
        "pending": len(catalog.PENDING),
        "pending_blocks": catalog.PENDING_BLOCKS,
        "preserved": len(preserved),
        "inputs_in_print_area": len(inputs),
        "inputs_with_value": [cell for cell in inputs if sheet[cell].value is not None],
        "unmapped_inputs": sorted(set(inputs) - mapped - pending, key=_cell_sort_key),
        "candidates": len(candidates),
        "candidates_confirmed": len(candidates & mapped),
        "candidates_pending": len(candidates & pending),
        "candidates_discarded": len(candidates - mapped - pending),
        "fields_outside_candidates": sorted(mapped - candidates, key=_cell_sort_key),
    }


def _cell_sort_key(coordinate):
    column, row, _, _ = range_boundaries(coordinate)[0:4]
    return (row, column)


def _candidate_cells(package_dir):
    """Células listadas no campos_para_revisar.json do pacote, quando existir."""
    path = Path(package_dir) / catalog.CANDIDATES_FILENAME
    if not path.is_file():
        return set()
    items = json.loads(path.read_text(encoding="utf-8"))
    return {
        item["cell_or_range"]
        for item in items
        if item.get("sheet_name") == catalog.SHEET_NAME and item.get("cell_or_range")
    }


def clean_copy(package_dir):
    """Cópia da planilha sem os dados do ensaio anterior nos campos variáveis."""
    raw = reference_bytes(package_dir)
    sheet = load_workbook(BytesIO(raw), data_only=False)[catalog.SHEET_NAME]
    preserved = preserved_coordinates(sheet)
    targets = [
        (catalog.SHEET_NAME, cell.coordinate)
        for cell in _cells_in(sheet, catalog.PRINT_AREA)
        if _is_input(cell, preserved)
    ]
    return build_cleaned_workbook(raw, targets)


@transaction.atomic
def register_tca_itmu(*, package_dir, version="1.0.0", activate=False, created_by=None):
    """Cadastra o modelo, a cópia limpa e os mapeamentos já confirmados."""
    report = audit(package_dir)
    if report["problems"]:
        raise PackageError("O catálogo não confere com a planilha de referência:\n- " + "\n- ".join(report["problems"]))
    if DocumentTemplate.objects.filter(code=catalog.TEMPLATE_CODE, version=version).exists():
        raise PackageError(
            "A versão {} do modelo {} já está cadastrada. Use outra versão; nada foi sobrescrito.".format(
                version, catalog.TEMPLATE_CODE
            )
        )

    workbook_bytes, cleared = clean_copy(package_dir)
    template = DocumentTemplate.objects.create(
        code=catalog.TEMPLATE_CODE,
        name=catalog.TEMPLATE_NAME,
        issuer=catalog.TEMPLATE_ISSUER,
        version=version,
        description=catalog.TEMPLATE_DESCRIPTION,
        page_count=1,
        repeat_layout_per_page=False,
        calibration_status="pending",
        is_active=activate,
        created_by=created_by,
    )

    fields = {}
    for order, item in enumerate(catalog.CONFIRMED, start=1):
        fields[item["field_id"]] = TemplateField.objects.create(
            template=template,
            field_id=item["field_id"],
            name=item["name"],
            section=item["section"],
            page_number=1,
            order=order,
            location_hint=item["anchor"],
            excel_sheet=catalog.SHEET_NAME,
            excel_cell=item["cell"],
            expected_type=item["expected_type"],
            validation_rule=item["rule"],
            required=item["required"],
        )
    # Campo em dúvida entra sem destino no Excel: fica visível para revisão e a
    # exportação não tem por onde escrever nele.
    for order, item in enumerate(catalog.PENDING, start=len(catalog.CONFIRMED) + 1):
        TemplateField.objects.create(
            template=template,
            field_id=item["field_id"],
            name=item["name"],
            section=item["section"],
            page_number=1,
            order=order,
            location_hint=item["anchor"],
            expected_type=item["expected_type"],
            validation_rule="CONFIRMAR ANTES DE MAPEAR ({}): {}".format(item["cell"], item["doubt"])[:220],
            required=False,
        )

    excel_template = ExcelTemplate.objects.create(
        document_template=template,
        name=EXCEL_TEMPLATE_NAME,
        version=version,
        workbook=ContentFile(workbook_bytes, name=CLEAN_WORKBOOK_NAME),
        is_active=activate,
    )
    for item in catalog.CONFIRMED:
        ExcelFieldMapping.objects.create(
            excel_template=excel_template,
            field=fields[item["field_id"]],
            sheet_name=catalog.SHEET_NAME,
            cell_or_range=item["cell"],
            expected_type=item["expected_type"],
        )

    report["template"] = template
    report["excel_template"] = excel_template
    report["cleared_cells"] = sorted(coordinate for _, coordinate in cleared)
    return report


def reviewed_candidates(package_dir):
    """Revisão do campos_para_revisar.json: o que virou campo e o que foi descartado."""
    path = Path(package_dir) / catalog.CANDIDATES_FILENAME
    if not path.is_file():
        raise PackageError("Arquivo {} não encontrado em {}.".format(catalog.CANDIDATES_FILENAME, package_dir))
    original = json.loads(path.read_text(encoding="utf-8"))
    sheet = load_workbook(BytesIO(reference_bytes(package_dir)), data_only=False)[catalog.SHEET_NAME]
    preserved = preserved_coordinates(sheet)
    confirmed = {item["cell"]: item for item in catalog.CONFIRMED}
    pending = {item["cell"]: item for item in catalog.PENDING}
    blocked = {}
    for block in catalog.PENDING_BLOCKS:
        for reference in block["ranges"]:
            for cell in _cells_in(sheet, reference):
                blocked[cell.coordinate] = block

    reviewed = []
    for candidate in original:
        cell = candidate.get("cell_or_range")
        entry = {
            "cell_or_range": cell,
            "sheet_name": candidate.get("sheet_name"),
            "original_field_id": candidate.get("field_id"),
            "example_value": candidate.get("example_value"),
        }
        if cell in confirmed:
            item = confirmed[cell]
            entry.update(
                status="campo_confirmado",
                approved=True,
                field_id=item["field_id"],
                name=item["name"],
                section=item["section"],
                expected_type=item["expected_type"],
                required=item["required"],
                location_hint=item["anchor"],
                note=item["rule"],
            )
        elif cell in pending:
            item = pending[cell]
            entry.update(
                status="precisa_confirmacao",
                approved=False,
                field_id=item["field_id"],
                name=item["name"],
                section=item["section"],
                expected_type=item["expected_type"],
                required=False,
                location_hint=item["anchor"],
                note=item["doubt"],
            )
        elif cell in preserved:
            entry.update(status="texto_fixo", approved=False, note=preserved[cell])
        elif cell in blocked:
            entry.update(status="bloco_em_revisao", approved=False, note=blocked[cell]["doubt"])
        else:
            entry.update(status="descartado", approved=False, note=_discard_reason(sheet, cell))
        reviewed.append(entry)

    extra = [
        {
            "cell_or_range": item["cell"],
            "sheet_name": catalog.SHEET_NAME,
            "original_field_id": None,
            "example_value": _as_text(sheet[item["cell"]].value),
            "status": "campo_confirmado",
            "approved": True,
            "field_id": item["field_id"],
            "name": item["name"],
            "section": item["section"],
            "expected_type": item["expected_type"],
            "required": item["required"],
            "location_hint": item["anchor"],
            "note": "âncora de célula mesclada: ficou fora da varredura original de candidatas",
        }
        for item in catalog.CONFIRMED
        if item["cell"] not in {candidate.get("cell_or_range") for candidate in original}
    ]
    return reviewed + extra


def _as_text(value):
    return None if value is None else str(value)


def _discard_reason(sheet, coordinate):
    try:
        cell = sheet[coordinate]
    except (TypeError, ValueError):
        return "endereço não localizado na planilha"
    min_col, min_row, max_col, max_row = range_boundaries(catalog.PRINT_AREA)
    if not (min_row <= cell.row <= max_row and min_col <= cell.column <= max_col):
        return "fora da área de impressão {}: coluna {} é tabela auxiliar ou margem".format(
            catalog.PRINT_AREA, get_column_letter(cell.column)
        )
    if cell.value is None:
        return "célula livre sem rótulo e vazia na referência: sobra estrutural da linha, não um campo"
    return "célula livre sem significado identificado na ficha impressa"

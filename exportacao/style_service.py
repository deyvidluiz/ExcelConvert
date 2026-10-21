"""Serviço central de estilos: função visual -> estilo openpyxl.

A cor da saída é decidida aqui e nunca depende de haver cor na origem. Um
documento fotografado em preto e branco recebe exatamente a mesma paleta que a
versão colorida da mesma ficha.

Prioridade, do mais forte ao mais fraco:
1. estilos de um modelo Excel de referência (``StyleConfig.from_excel_template``);
2. indícios de estilo observados com confiança (tamanho/negrito/itálico apenas,
   combinados com a política de cores configurada);
3. paleta padrão configurável.
"""

import json
import os
from copy import copy
from io import BytesIO

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from extracao.schemas import CellFunction


DEFAULT_PALETTE = {
    "title": {"background": "1F4E79", "foreground": "FFFFFF"},
    "header": {"background": "3C5CCF", "foreground": "FFFFFF"},
    "subheader": {"background": "9DC3E6", "foreground": "182230"},
    "data": {"background": "FFFFFF", "foreground": "000000"},
    "review": {"background": "FFF2CC", "foreground": "7F5F00"},
    "footer": {"background": "E7E6E6", "foreground": "000000"},
    "empty": {"background": "FFFFFF", "foreground": "000000"},
}

EVIDENCE_THRESHOLD = 0.75
EMPHASIS_ROLES = {CellFunction.TITLE.value, CellFunction.HEADER.value}
BOLD_ROLES = EMPHASIS_ROLES | {CellFunction.SUBHEADER.value}
BORDER_STYLES = {"thin", "medium", "thick", "hair", "dotted", "dashed", "double"}


def _clean_colour(value):
    if not value:
        return None
    text = str(value).strip().lstrip("#").upper()
    if len(text) == 8:
        text = text[2:]
    if len(text) == 6 and all(char in "0123456789ABCDEF" for char in text):
        return text
    return None


def configured_palette():
    """Paleta padrão: settings.EXCEL_PALETTE, sobreposta por EXCEL_PALETTE_JSON."""
    palette = {}
    try:
        from django.conf import settings

        palette = dict(getattr(settings, "EXCEL_PALETTE", {}) or {})
    except Exception:
        palette = {}
    raw = os.getenv("EXCEL_PALETTE_JSON", "").strip()
    if raw:
        try:
            override = json.loads(raw)
        except (TypeError, ValueError):
            override = {}
        if isinstance(override, dict):
            for role, colours in override.items():
                if isinstance(colours, dict):
                    palette[role] = {**palette.get(role, {}), **colours}
    return palette


class StyleConfig:
    """Política de cores independente das cores presentes no documento."""

    BORDER_COLOR = "B7C9D6"
    ACCENT_BORDER_COLOR = "3C5CCF"

    def __init__(self, palette=None, *, source="default"):
        merged = {}
        for role, default in DEFAULT_PALETTE.items():
            override = (palette or {}).get(role) or {}
            merged[role] = {
                "background": _clean_colour(override.get("background")) or default["background"],
                "foreground": _clean_colour(override.get("foreground")) or default["foreground"],
            }
        self.palette = merged
        self.source = source

    @classmethod
    def from_settings(cls):
        return cls(configured_palette())

    @classmethod
    def from_excel_template(cls, excel_template):
        """Deriva a paleta do modelo Excel cadastrado, caindo para a padrão."""
        try:
            palette = palette_from_workbook(excel_template.workbook.read())
        except Exception:
            palette = None
        finally:
            try:
                excel_template.workbook.seek(0)
            except Exception:
                pass
        if not palette:
            return cls.from_settings()
        return cls({**configured_palette(), **palette}, source="excel_template")

    def colours_for(self, role):
        return self.palette.get(role, self.palette[CellFunction.DATA.value])

    def style_for(self, function, *, alignment=None, border=None, evidence=None, has_borders=True):
        role = getattr(function, "value", function) or CellFunction.DATA.value
        if role not in self.palette:
            role = CellFunction.DATA.value
        colours = self.colours_for(role)

        font_kwargs = {"color": colours["foreground"], "size": 10, "name": "Calibri"}
        if role in BOLD_ROLES:
            font_kwargs["bold"] = True
        if evidence and (evidence.get("confidence") or 0) >= EVIDENCE_THRESHOLD:
            if evidence.get("font_size"):
                font_kwargs["size"] = min(28, max(7, round(float(evidence["font_size"]), 1)))
            if evidence.get("bold") is not None:
                font_kwargs["bold"] = bool(evidence["bold"]) or role in BOLD_ROLES
            if evidence.get("italic") is not None:
                font_kwargs["italic"] = bool(evidence["italic"])

        horizontal = alignment if alignment in {"left", "center", "right", "general"} else None
        if horizontal is None:
            horizontal = "center" if role in EMPHASIS_ROLES else "left"

        return {
            "font": Font(**font_kwargs),
            "fill": PatternFill("solid", fgColor=colours["background"]),
            "alignment": Alignment(horizontal=horizontal, vertical="center", wrap_text=True),
            "border": self._border_for(role, border, has_borders),
        }

    def _border_for(self, role, border, has_borders):
        emphasis = role in EMPHASIS_ROLES
        if not has_borders and not emphasis and role != CellFunction.SUBHEADER.value and not border:
            return Border()
        default_style = "medium" if emphasis else "thin"
        colour = self.ACCENT_BORDER_COLOR if emphasis else self.BORDER_COLOR
        sides = {name: Side(style=default_style, color=colour) for name in ("top", "right", "bottom", "left")}
        if border:
            for name in sides:
                value = border.get(name)
                sides[name] = Side(style=value, color=colour) if value in BORDER_STYLES else Side()
        return Border(**sides)


def palette_from_workbook(raw_workbook):
    """Infere papéis visuais a partir dos preenchimentos usados em um modelo."""
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(raw_workbook), data_only=True)
    scored = {}
    for sheet in workbook.worksheets[:3]:
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 1, 60)):
            for cell in row:
                fill = cell.fill
                if not fill or fill.patternType != "solid":
                    continue
                background = _clean_colour(getattr(fill.fgColor, "rgb", None))
                if not background or background == "FFFFFF":
                    continue
                foreground = _clean_colour(getattr(cell.font.color, "rgb", None)) or "000000"
                entry = scored.setdefault(background, {"count": 0, "foreground": foreground, "row": cell.row})
                entry["count"] += 1
                entry["row"] = min(entry["row"], cell.row)
    if not scored:
        return None

    def luminance(colour):
        red, green, blue = (int(colour[index : index + 2], 16) for index in (0, 2, 4))
        return 0.2126 * red + 0.7152 * green + 0.0722 * blue

    ordered = sorted(scored.items(), key=lambda item: luminance(item[0]))
    palette = {}
    darkest, dark_entry = ordered[0]
    palette["title"] = {"background": darkest, "foreground": dark_entry["foreground"]}
    header, header_entry = ordered[1] if len(ordered) > 1 else (darkest, dark_entry)
    palette["header"] = {"background": header, "foreground": header_entry["foreground"]}
    lightest, light_entry = ordered[-1]
    palette["subheader"] = {"background": lightest, "foreground": light_entry["foreground"]}
    return palette


def get_style_for_function(function, is_header=False, *, config=None, metadata=None):
    """API compatível com o exportador antigo, agora centralizada."""
    config = config or StyleConfig.from_settings()
    metadata = metadata or {}
    return config.style_for(
        function or (CellFunction.HEADER.value if is_header else CellFunction.DATA.value),
        alignment=metadata.get("alignment"),
        border=metadata.get("border"),
        evidence=metadata.get("style_evidence"),
        has_borders=metadata.get("has_borders", True),
    )


def apply_cell_style(cell, function, *, config=None, metadata=None):
    for attribute, value in get_style_for_function(function, config=config, metadata=metadata).items():
        setattr(cell, attribute, copy(value))

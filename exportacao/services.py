import posixpath
import re
import zipfile
from html import escape
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree

from django.conf import settings
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import range_boundaries
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from docx import Document as WordDocument
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor

from extracao.schemas import CellFunction, normalise_layout_payload

from .layout_export import build_xlsx_with_layout
from .style_service import StyleConfig


def extraction_for(document):
    if document.extraction_payload:
        return document.extraction_payload
    from documentos.models import ExtractionCache

    cache = ExtractionCache.objects.filter(sha256=document.sha256).order_by("-created_at").first()
    if not cache:
        raise ValueError("Os dados ainda não foram processados.")
    return cache.payload


BOLD_FUNCTIONS = {CellFunction.TITLE.value, CellFunction.HEADER.value, CellFunction.SUBHEADER.value}
CENTERED_FUNCTIONS = {CellFunction.TITLE.value, CellFunction.HEADER.value}


def _layout_pages(payload):
    """Normaliza o payload para que Word e PDF leiam o mesmo contrato do XLSX."""
    return normalise_layout_payload(payload).get("pages") or []


def _is_landscape(pages):
    for page in pages:
        width = page.get("page_width") or 0
        height = page.get("page_height") or 0
        if width and height:
            return width > height
    # Sem geometria, mantém a paisagem usada pelas versões anteriores.
    return True


def _split_regions(page):
    regions = [region for region in page.get("regions") or [] if (region.get("value") or "").strip()]
    leading = [region for region in regions if region.get("function") != CellFunction.FOOTER.value]
    trailing = [region for region in regions if region.get("function") == CellFunction.FOOTER.value]
    return leading, trailing


def _span(value):
    try:
        return max(1, int(value))
    except (TypeError, ValueError):
        return 1


def _spanned_matrix(table):
    """Distribui as células numa matriz, reservando as áreas mescladas."""
    lines = []
    if table.get("headers"):
        lines.append((CellFunction.HEADER.value, table["headers"]))
    for row in table.get("rows") or []:
        lines.append((CellFunction.DATA.value, row))

    placements = []
    occupied = set()
    columns = 1
    rows = len(lines)
    for row_index, (default, line) in enumerate(lines):
        col_index = 0
        for raw in line:
            cell = dict(raw) if isinstance(raw, dict) else {"value": raw}
            cell["function"] = cell.get("function") or default
            while (row_index, col_index) in occupied:
                col_index += 1
            row_span = _span(cell.get("row_span"))
            col_span = _span(cell.get("col_span"))
            for row_offset in range(row_span):
                for col_offset in range(col_span):
                    occupied.add((row_index + row_offset, col_index + col_offset))
            placements.append((row_index, col_index, row_span, col_span, cell))
            columns = max(columns, col_index + col_span)
            rows = max(rows, row_index + row_span)
            col_index += col_span

    if not placements:
        return [], [], columns

    matrix = [[None] * columns for _ in range(rows)]
    spans = []
    for row_index, col_index, row_span, col_span, cell in placements:
        matrix[row_index][col_index] = cell
        if row_span > 1 or col_span > 1:
            spans.append((row_index, col_index, row_index + row_span - 1, col_index + col_span - 1))
    return matrix, spans, columns


def _column_shares(table, columns):
    """Proporção de largura por coluna, vinda da geometria quando existir."""
    values = table.get("column_widths") or table.get("column_widths_relative") or []
    numbers = [max(0.0, float(value or 0)) for value in list(values)[:columns]]
    numbers += [0.0] * max(0, columns - len(numbers))
    total = sum(numbers)
    if not total:
        return [1 / columns] * columns
    numbers = [value or total / columns for value in numbers]
    total = sum(numbers)
    return [value / total for value in numbers]


def build_xlsx(payload, *, style_config=None):
    """Exporta o payload estruturado; payloads antigos recebem layout compatível."""
    return build_xlsx_with_layout(payload, config=style_config)


def style_config_for(document):
    """Prioriza os estilos do modelo Excel associado ao tipo de ficha."""
    if getattr(document, "template_id", None):
        excel_template = document.template.excel_templates.filter(is_active=True).order_by("-created_at").first()
        if excel_template:
            return StyleConfig.from_excel_template(excel_template)
    return StyleConfig.from_settings()


def build_xlsx_from_template(payload, excel_template):
    """Preenche um modelo cadastrado preservando estrutura, estilos e fórmulas."""
    raw_template = excel_template.workbook.read()
    keep_vba = Path(excel_template.workbook.name or "").suffix.lower() == ".xlsm"
    workbook = _load_template_workbook(raw_template, keep_vba=keep_vba)
    mappings = {mapping.field.field_id: mapping for mapping in excel_template.mappings.select_related("field")}
    touched_cells = set()
    extracted = {}
    for field in payload.get("template_fields", []):
        mapping = mappings.get(field.get("field_id"))
        if not mapping or mapping.sheet_name not in workbook.sheetnames:
            continue
        extracted[field.get("field_id")] = True
        cells = _mapping_cells(workbook[mapping.sheet_name], mapping.cell_or_range)
        values = _values_for_cells(field.get("value"), len(cells), field.get("field_id"))
        for cell, value in zip(cells, values):
            if cell.data_type == "f":
                raise ValueError("O mapeamento {} aponta para uma fórmula em {}.".format(field.get("field_id"), cell.coordinate))
            cell.value = _coerce_expected_type(value, mapping.expected_type)
            touched_cells.add((mapping.sheet_name, cell.coordinate))
    touched_cells |= _clear_unextracted_mappings(workbook, mappings, extracted)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    preserved_output = _preserve_template_parts(_decrypt_template_bytes(raw_template), output.getvalue())
    _validate_saved_template(raw_template, preserved_output, keep_vba=keep_vba, touched_cells=touched_cells)
    return BytesIO(preserved_output)


NUMERIC_EXPECTED_TYPES = {"number", "decimal"}
# Inteiro, decimal com um único separador, ou milhar com ponto e decimal com vírgula.
_NUMERIC_TEXT = re.compile(r"-?\d{1,3}(?:\.\d{3})+,\d+|-?\d+[.,]\d+|-?\d+")


def _coerce_expected_type(value, expected_type):
    """Campo numérico vira número; o resto é gravado exatamente como foi lido.

    As fórmulas da ficha comparam e procuram valores numéricos: um "10" ou um
    "31,6" gravados como texto fazem `K44=10` dar falso e o `VLOOKUP` da correção
    de temperatura cair na última faixa sem avisar. Só converte o que é
    inequivocamente numérico: ">5000000", "300x600 - 5A" e datas seguem texto.
    """
    if expected_type not in NUMERIC_EXPECTED_TYPES or not isinstance(value, str):
        return value
    text = value.strip()
    if not _NUMERIC_TEXT.fullmatch(text):
        return value
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    return float(text) if "." in text else int(text)


def _clear_unextracted_mappings(workbook, mappings, extracted):
    """Campo sem valor na extração não pode herdar o dado do arquivo-base."""
    cleared = set()
    for field_id, mapping in mappings.items():
        if field_id in extracted or mapping.sheet_name not in workbook.sheetnames:
            continue
        for cell in _mapping_cells(workbook[mapping.sheet_name], mapping.cell_or_range):
            if cell.data_type == "f" or cell.value is None:
                continue
            cell.value = None
            cleared.add((mapping.sheet_name, cell.coordinate))
    return cleared


def build_cleaned_workbook(raw_template, targets):
    """Copia o modelo esvaziando apenas os destinos variáveis informados.

    Rótulos, títulos, fórmulas, estilos, mesclagens, dimensões, área de impressão
    e tabelas auxiliares seguem intactos: só as células recebidas ficam vazias.
    """
    decrypted = _decrypt_template_bytes(raw_template)
    workbook = load_workbook(BytesIO(decrypted), data_only=False)
    cleared = set()
    for sheet_name, coordinate in targets:
        if sheet_name not in workbook.sheetnames:
            raise ValueError("O modelo não tem a aba {}.".format(sheet_name))
        cell = workbook[sheet_name][coordinate]
        if isinstance(cell, MergedCell):
            raise ValueError("{}!{} é célula mesclada secundária.".format(sheet_name, coordinate))
        if cell.data_type == "f":
            raise ValueError("{}!{} contém fórmula e não pode ser esvaziada.".format(sheet_name, coordinate))
        if cell.value is None:
            continue
        cell.value = None
        cleared.add((sheet_name, cell.coordinate))
    output = BytesIO()
    workbook.save(output)
    preserved = _preserve_template_parts(decrypted, output.getvalue())
    _validate_saved_template(decrypted, preserved, keep_vba=False, touched_cells=cleared)
    return preserved, cleared


def _decrypt_template_bytes(raw_template):
    if raw_template.startswith(b"PK" + bytes((3, 4))):
        return raw_template
    if not raw_template.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        raise ValueError("O arquivo do modelo não é um XLSX/XLSM reconhecido.")
    password = getattr(settings, "EXCEL_TEMPLATE_PASSWORD", "")
    if not password:
        raise ValueError("EXCEL_TEMPLATE_PASSWORD não está configurada para o modelo Excel protegido.")
    try:
        import msoffcrypto
    except ImportError as exc:
        raise ValueError("A dependência msoffcrypto-tool é necessária para abrir o modelo Excel protegido.") from exc
    decrypted = BytesIO()
    try:
        office = msoffcrypto.OfficeFile(BytesIO(raw_template))
        office.load_key(password=password)
        office.decrypt(decrypted)
    except Exception as exc:
        raise ValueError("Não foi possível abrir o modelo Excel protegido com a senha configurada.") from exc
    return decrypted.getvalue()


def _load_template_workbook(raw_template, *, keep_vba=False):
    """Abre XLSX normal ou pacote Office cifrado sem expor a senha."""
    return load_workbook(BytesIO(_decrypt_template_bytes(raw_template)), keep_vba=keep_vba, data_only=False)


def _mapping_cells(sheet, cell_or_range):
    """Expande células e intervalos, aceitando grupos separados por ';'."""
    addresses = [part.strip() for part in re.split(r"[;,]", cell_or_range or "") if part.strip()]
    cells = []
    for address in addresses:
        try:
            min_col, min_row, max_col, max_row = range_boundaries(address)
        except ValueError as exc:
            raise ValueError("Destino Excel inválido: {}".format(address)) from exc
        for row in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                if isinstance(cell, MergedCell):
                    raise ValueError(
                        "Destino Excel não pode ser uma célula mesclada secundária: {}!{}".format(
                            sheet.title, cell.coordinate
                        )
                    )
                cells.append(cell)
    if not cells:
        raise ValueError("Destino Excel vazio: {}".format(cell_or_range))
    return cells


def _values_for_cells(value, cell_count, field_id):
    if isinstance(value, dict):
        value = value.get("values", value.get("value"))
    if isinstance(value, (list, tuple)):
        if len(value) != cell_count:
            raise ValueError("O campo {} exige {} valores; foram recebidos {}.".format(field_id, cell_count, len(value)))
        return list(value)
    if cell_count != 1:
        raise ValueError(
            "O campo {} está mapeado para {} células e precisa de uma lista estruturada.".format(field_id, cell_count)
        )
    return [value]


def _preserve_template_parts(raw_template, output):
    """Restaura partes OOXML que o openpyxl não modela, como impressão e metadados."""
    source_zip = zipfile.ZipFile(BytesIO(raw_template))
    output_zip = zipfile.ZipFile(BytesIO(output))
    source_names = set(source_zip.namelist())
    output_parts = {name: output_zip.read(name) for name in output_zip.namelist()}
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_attr = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    ElementTree.register_namespace("", main_ns)
    ElementTree.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")

    content_types_name = "[Content_Types].xml"
    content_types = ElementTree.fromstring(output_parts[content_types_name])
    existing_overrides = {item.attrib.get("PartName") for item in content_types}
    for name in sorted(source_names - set(output_parts)):
        if name.startswith(("xl/printerSettings/", "xl/customProperty", "docMetadata/")):
            output_parts[name] = source_zip.read(name)
            source_content = ElementTree.fromstring(source_zip.read(content_types_name))
            for item in source_content:
                part_name = item.attrib.get("PartName", "")
                if part_name.lstrip("/") == name and part_name not in existing_overrides:
                    content_types.append(item)
                    existing_overrides.add(part_name)

    for rels_name in sorted(name for name in source_names if name.startswith("xl/worksheets/_rels/") and name.endswith(".rels")):
        source_rels = ElementTree.fromstring(source_zip.read(rels_name))
        output_rels_name = rels_name
        if output_rels_name not in output_parts:
            continue
        output_rels = ElementTree.fromstring(output_parts[output_rels_name])
        used_ids = {item.attrib.get("Id") for item in output_rels}
        next_id = 1
        replacements = {}
        source_dir = posixpath.dirname(posixpath.dirname(rels_name))
        for relationship in source_rels:
            rel_type = relationship.attrib.get("Type", "")
            if not rel_type.endswith(("/printerSettings", "/customProperty")):
                continue
            target = relationship.attrib.get("Target", "")
            source_part = posixpath.normpath(posixpath.join(source_dir, target))
            if source_part not in output_parts and source_part in source_names:
                output_parts[source_part] = source_zip.read(source_part)
            if source_part not in output_parts:
                continue
            while "rId{}".format(next_id) in used_ids:
                next_id += 1
            new_id = "rId{}".format(next_id)
            next_id += 1
            copied = dict(relationship.attrib)
            copied["Id"] = new_id
            output_rels.append(ElementTree.Element(relationship.tag, copied))
            used_ids.add(new_id)
            replacements[relationship.attrib.get("Id")] = new_id

        sheet_name = rels_name.replace("/_rels/", "/").removesuffix(".rels")
        if sheet_name in output_parts and replacements:
            source_sheet_xml = ElementTree.fromstring(source_zip.read(sheet_name))
            source_page_setup_ids = [
                page_setup.attrib.get(rel_attr)
                for page_setup in source_sheet_xml.iter("{" + main_ns + "}pageSetup")
                if page_setup.attrib.get(rel_attr)
            ]
            sheet_xml = ElementTree.fromstring(output_parts[sheet_name])
            for page_setup in sheet_xml.iter("{" + main_ns + "}pageSetup"):
                current_id = page_setup.attrib.get(rel_attr)
                original_id = current_id or (source_page_setup_ids[0] if source_page_setup_ids else None)
                if original_id in replacements:
                    page_setup.attrib[rel_attr] = replacements[original_id]
            output_parts[sheet_name] = ElementTree.tostring(sheet_xml, encoding="utf-8", xml_declaration=True)
        output_parts[output_rels_name] = ElementTree.tostring(output_rels, encoding="utf-8", xml_declaration=True)

    output_parts[content_types_name] = ElementTree.tostring(content_types, encoding="utf-8", xml_declaration=True)
    restored = BytesIO()
    with zipfile.ZipFile(restored, "w", zipfile.ZIP_DEFLATED) as destination:
        for name, data in output_parts.items():
            destination.writestr(name, data)
    return restored.getvalue()


def _workbook_signature(workbook, *, touched=()):
    """Assina estrutura e fórmulas. Os destinos tocados saem da conta da dimensão."""
    for sheet_name, coordinate in touched:
        if sheet_name in workbook.sheetnames:
            cell = workbook[sheet_name][coordinate]
            if not isinstance(cell, MergedCell):
                cell.value = None
    formulas = {}
    sheets = {}
    for sheet in workbook.worksheets:
        sheets[sheet.title] = {
            "dimension": sheet.calculate_dimension(),
            "merged": sorted(str(item) for item in sheet.merged_cells.ranges),
            "print_area": str(sheet.print_area),
            "orientation": sheet.page_setup.orientation,
            "paper_size": sheet.page_setup.paperSize,
            "column_widths": sorted((key, value.width, value.hidden) for key, value in sheet.column_dimensions.items()),
            "row_heights": sorted((key, value.height, value.hidden) for key, value in sheet.row_dimensions.items()),
            "images": len(sheet._images),
            "charts": len(sheet._charts),
        }
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formulas[(sheet.title, cell.coordinate)] = cell.value
    return {"sheetnames": list(workbook.sheetnames), "sheets": sheets, "formulas": formulas}


def _validate_saved_template(raw_template, output, *, keep_vba, touched_cells):
    before = _workbook_signature(_load_template_workbook(raw_template, keep_vba=keep_vba), touched=touched_cells)
    after_workbook = load_workbook(BytesIO(output), keep_vba=keep_vba, data_only=False)
    after = _workbook_signature(after_workbook, touched=touched_cells)
    if before["sheetnames"] != after["sheetnames"] or before["sheets"] != after["sheets"]:
        raise ValueError("A exportação alterou a estrutura, mesclagens, impressão ou dimensões do modelo.")
    if before["formulas"] != after["formulas"]:
        raise ValueError("A exportação alterou fórmulas existentes no modelo.")


def build_docx(payload, *, style_config=None):
    """Gera o Word com a mesma paleta e os mesmos papéis visuais do XLSX."""
    config = style_config or StyleConfig.from_settings()
    pages = _layout_pages(payload)
    document = WordDocument()
    section = document.sections[0]
    if _is_landscape(pages):
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    usable = section.page_width - section.left_margin - section.right_margin

    written = False
    for page in pages:
        if written:
            document.add_page_break()
        leading, trailing = _split_regions(page)
        written = _docx_regions(document, leading, config) or written
        for table in page.get("tables") or []:
            if table.get("title"):
                _docx_paragraph(document, table["title"], CellFunction.TITLE.value, config)
            written = _docx_table(document, table, config, usable) or written
        written = _docx_regions(document, trailing, config) or written

    if not written:
        _docx_paragraph(document, "Nenhuma tabela ou texto foi identificado no documento.", CellFunction.REVIEW.value, config)
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def _docx_regions(document, regions, config):
    for region in regions:
        _docx_paragraph(
            document,
            region.get("value") or "",
            region.get("function") or CellFunction.DATA.value,
            config,
            alignment=region.get("alignment"),
        )
    return bool(regions)


def _docx_paragraph(document, text, function, config, *, alignment=None):
    paragraph = document.add_paragraph()
    paragraph.alignment = _docx_alignment(function, alignment)
    _docx_write_runs(paragraph, text, function, config)
    colours = config.colours_for(function)
    if colours["background"] != "FFFFFF":
        _shade(paragraph._p.get_or_add_pPr(), colours["background"])
    return paragraph


def _docx_table(document, table, config, usable):
    matrix, spans, width = _spanned_matrix(table)
    if not matrix:
        return False
    word_table = document.add_table(rows=len(matrix), cols=width)
    word_table.style = "Table Grid" if table.get("has_borders", True) else "Normal Table"
    word_table.autofit = False

    for row_index, line in enumerate(matrix):
        for col_index, cell in enumerate(line):
            target = word_table.cell(row_index, col_index)
            function = (cell or {}).get("function") or CellFunction.EMPTY.value
            colours = config.colours_for(function)
            paragraph = target.paragraphs[0]
            paragraph.alignment = _docx_alignment(function, (cell or {}).get("alignment"))
            _docx_write_runs(paragraph, (cell or {}).get("value") or "", function, config)
            if colours["background"] != "FFFFFF":
                _shade(target._tc.get_or_add_tcPr(), colours["background"])

    # As larguras vêm antes da mesclagem: depois, a célula mesclada cobre várias colunas.
    for col_index, share in enumerate(_column_shares(table, width)):
        size = Emu(int(usable * share))
        word_table.columns[col_index].width = size
        for row in word_table.rows:
            if col_index < len(row.cells):
                row.cells[col_index].width = size

    for start_row, start_col, end_row, end_col in spans:
        word_table.cell(start_row, start_col).merge(word_table.cell(end_row, end_col))
    return True


def _docx_write_runs(paragraph, text, function, config):
    """Preserva quebras de linha do valor original em um único parágrafo."""
    colours = config.colours_for(function)
    lines = str(text or "").split("\n")
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.size = Pt(9)
        run.font.bold = function in BOLD_FUNCTIONS
        run.font.color.rgb = RGBColor.from_string(colours["foreground"])
        if index < len(lines) - 1:
            run.add_break()


def _docx_alignment(function, alignment):
    horizontal = alignment if alignment in {"left", "center", "right"} else None
    if horizontal is None:
        horizontal = "center" if function in CENTERED_FUNCTIONS else "left"
    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[horizontal]


def _shade(properties, colour):
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:color"), "auto")
    element.set(qn("w:fill"), colour)
    properties.append(element)


def build_pdf(payload, *, style_config=None):
    """Gera o PDF com a mesma paleta e os mesmos papéis visuais do XLSX."""
    config = style_config or StyleConfig.from_settings()
    pages = _layout_pages(payload)
    output = BytesIO()
    pagesize = landscape(A4) if _is_landscape(pages) else A4
    doc = SimpleDocTemplate(
        output,
        pagesize=pagesize,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    story = []
    for page in pages:
        if story:
            story.append(PageBreak())
        leading, trailing = _split_regions(page)
        story.extend(_pdf_regions(leading, config, doc.width))
        for table in page.get("tables") or []:
            if table.get("title"):
                story.extend(_pdf_regions([{"value": table["title"], "function": CellFunction.TITLE.value}], config, doc.width))
            flowable = _pdf_table(table, config, doc.width)
            if flowable is not None:
                story.extend([flowable, Spacer(1, 4 * mm)])
        story.extend(_pdf_regions(trailing, config, doc.width))

    if not story:
        story.append(Paragraph("Nenhuma tabela ou texto foi identificado no documento.", _pdf_style(CellFunction.REVIEW.value, config)))
    doc.build(story)
    output.seek(0)
    return output


def _pdf_regions(regions, config, width):
    """Cada texto solto vira uma faixa com o fundo do seu papel visual."""
    flowables = []
    for region in regions:
        function = region.get("function") or CellFunction.DATA.value
        style = _pdf_style(function, config, alignment=region.get("alignment"))
        colours = config.colours_for(function)
        paragraph = Paragraph(_pdf_text(region.get("value")), style)
        if colours["background"] == "FFFFFF":
            flowables.extend([paragraph, Spacer(1, 2 * mm)])
            continue
        banner = Table([[paragraph]], colWidths=[width])
        banner.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#" + colours["background"])),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        flowables.extend([banner, Spacer(1, 2 * mm)])
    return flowables


def _pdf_table(table, config, width):
    matrix, spans, columns = _spanned_matrix(table)
    if not matrix:
        return None
    data = []
    commands = [("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
    for row_index, line in enumerate(matrix):
        row = []
        for col_index, cell in enumerate(line):
            function = (cell or {}).get("function") or CellFunction.EMPTY.value
            colours = config.colours_for(function)
            row.append(Paragraph(_pdf_text((cell or {}).get("value")), _pdf_style(function, config, alignment=(cell or {}).get("alignment"))))
            if colours["background"] != "FFFFFF":
                commands.append(("BACKGROUND", (col_index, row_index), (col_index, row_index), colors.HexColor("#" + colours["background"])))
        data.append(row)

    for start_row, start_col, end_row, end_col in spans:
        commands.append(("SPAN", (start_col, start_row), (end_col, end_row)))

    if table.get("has_borders", True):
        commands.append(("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#" + StyleConfig.BORDER_COLOR)))
    elif table.get("headers"):
        commands.append(("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor("#" + StyleConfig.ACCENT_BORDER_COLOR)))

    flowable = Table(
        data,
        colWidths=[width * share for share in _column_shares(table, columns)],
        repeatRows=1 if table.get("headers") else 0,
    )
    flowable.setStyle(TableStyle(commands))
    return flowable


def _pdf_style(function, config, *, alignment=None):
    colours = config.colours_for(function)
    horizontal = alignment if alignment in {"left", "center", "right"} else None
    if horizontal is None:
        horizontal = "center" if function in CENTERED_FUNCTIONS else "left"
    return ParagraphStyle(
        "cell-{}-{}".format(function, horizontal),
        fontName="Helvetica-Bold" if function in BOLD_FUNCTIONS else "Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#" + colours["foreground"]),
        alignment={"left": TA_LEFT, "center": TA_CENTER, "right": TA_RIGHT}[horizontal],
    )


def _pdf_text(value):
    return escape(str(value if value is not None else "")).replace("\n", "<br/>")

"""Exportação XLSX que respeita a estrutura reconstruída do documento."""

from copy import copy
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from extracao.schemas import CellFunction, normalise_layout_payload

from .layout_grid import plan_page
from .style_service import StyleConfig, apply_cell_style

MAX_SHEET_TITLE = 31


def build_xlsx_with_layout(payload: dict, *, palette=None, config=None):
    """Gera a planilha editável preservando a organização da folha original."""
    payload = normalise_layout_payload(payload)
    if config is None:
        config = StyleConfig(palette) if palette is not None else StyleConfig.from_settings()

    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles = set()

    for index, page in enumerate(payload.get("pages", []), start=1):
        plan = plan_page(page, index)
        if not plan.cells:
            continue
        sheet = workbook.create_sheet(title=_unique_title(plan.title, used_titles))
        _write_plan(sheet, plan, config)

    if not workbook.worksheets:
        sheet = workbook.create_sheet("Sem dados")
        sheet["A1"] = "Nenhuma tabela ou texto foi identificado no documento."
        apply_cell_style(sheet["A1"], CellFunction.REVIEW.value, config=config)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _unique_title(title, used_titles):
    base = (title or "Página")[:MAX_SHEET_TITLE]
    candidate = base
    suffix = 2
    while candidate in used_titles:
        tail = f" ({suffix})"
        candidate = base[: MAX_SHEET_TITLE - len(tail)] + tail
        suffix += 1
    used_titles.add(candidate)
    return candidate


def _write_plan(sheet, plan, config):
    for placed in plan.cells:
        target = sheet.cell(placed.row, placed.col)
        if isinstance(target, MergedCell):
            continue
        value = placed.value
        if value is not None:
            target.value = value
            if isinstance(value, str):
                target.number_format = "@"
                target.quotePrefix = True
        apply_cell_style(target, placed.function, config=config, metadata=placed.metadata)

    _apply_merges(sheet, plan.merges)

    for column, width in plan.column_widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = round(width, 2)
    for row, height in plan.row_heights.items():
        sheet.row_dimensions[row].height = round(height, 2)

    if plan.freeze_panes:
        sheet.freeze_panes = plan.freeze_panes
    _apply_sheet_settings(sheet, plan)


def _apply_merges(sheet, merges):
    claimed = set()
    for merged in merges:
        bounds = _merge_bounds(merged)
        if not bounds:
            continue
        start_row, start_col, end_row, end_col = bounds
        if end_row < start_row or end_col < start_col:
            continue
        if start_row == end_row and start_col == end_col:
            continue
        coordinates = {
            (row, col)
            for row in range(start_row, end_row + 1)
            for col in range(start_col, end_col + 1)
        }
        if coordinates & claimed:
            continue
        anchor = sheet.cell(start_row, start_col)
        style = (anchor.font, anchor.fill, anchor.border, anchor.alignment)
        try:
            sheet.merge_cells(start_row=start_row, start_column=start_col, end_row=end_row, end_column=end_col)
        except (TypeError, ValueError):
            continue
        # O Excel pinta a área mesclada com o estilo de cada célula de base.
        for row, col in coordinates:
            cell = sheet.cell(row, col)
            cell.font, cell.fill, cell.border, cell.alignment = (copy(item) for item in style)
        claimed |= coordinates


def _merge_bounds(merged):
    if isinstance(merged, str):
        try:
            min_col, min_row, max_col, max_row = range_boundaries(merged)
        except (TypeError, ValueError):
            return None
        return min_row, min_col, max_row, max_col
    try:
        start_row, start_col, end_row, end_col = (int(item) for item in merged)
    except (TypeError, ValueError):
        return None
    if start_row < 1 or start_col < 1:
        return None
    return start_row, start_col, end_row, end_col


def _apply_sheet_settings(sheet, plan):
    sheet.page_setup.orientation = "landscape" if plan.landscape else "portrait"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.35
    sheet.page_margins.bottom = 0.35


def estimate_xlsx_size(payload):
    return sum(
        len(table.get("headers", [])) + sum(len(row) for row in table.get("rows", []))
        for page in payload.get("pages", [])
        for table in page.get("tables", [])
    ) * 120

from io import BytesIO

from openpyxl import load_workbook

from extracao.local_pdf import extract_pdf_tables
from extracao.schemas import LAYOUT_SCHEMA_VERSION, normalise_layout_payload
from exportacao.services import build_xlsx


def _payload(values):
    return {
        "pages": [{
            "page_number": 1,
            "page_width": 600,
            "page_height": 800,
            "tables": [{
                "title": "Ficha de ensaio",
                "headers": [
                    {"value": "Campo", "function": "header"},
                    {"value": "Valor", "function": "header"},
                ],
                "rows": [
                    [{"value": "Seção", "function": "subheader"}, {"value": values[0], "function": "data"}],
                    [{"value": "Número", "function": "data"}, {"value": values[1], "function": "review", "ambiguous": True}],
                    [{"value": "Observação longa", "function": "data", "col_span": 2}],
                ],
                "column_widths": [80, 240],
                "merged_cells": [],
            }],
            "alerts": [],
        }],
        "alerts": [],
    }


def test_cores_e_organizacao_independem_da_origem_mono_ou_colorida():
    color = load_workbook(build_xlsx(_payload(("azul", "12")))).active
    mono = load_workbook(build_xlsx(_payload(("azul", "12")))).active

    assert [color.cell(1, 1).value, color.cell(3, 1).value, color.cell(4, 2).value] == [
        mono.cell(1, 1).value, mono.cell(3, 1).value, mono.cell(4, 2).value
    ]
    assert color["A1"].fill.fgColor.rgb.endswith("1F4E79")
    assert color["A3"].fill.fgColor.rgb.endswith("3C5CCF")
    assert color["A4"].fill.fgColor.rgb.endswith("9DC3E6")
    assert color["B5"].fill.fgColor.rgb.endswith("FFF2CC")
    assert "A6:B6" in {str(item) for item in color.merged_cells.ranges}
    assert color.column_dimensions["A"].width != color.column_dimensions["B"].width


def test_texto_numeros_e_metadados_de_layout_sao_preservados():
    payload = normalise_layout_payload(_payload(("10,50", "0012")))
    payload["pages"][0]["tables"][0]["rows"][1][1]["value"] = "0013"
    workbook = load_workbook(build_xlsx(payload), data_only=False)
    sheet = workbook.active

    assert payload["layout_version"] == LAYOUT_SCHEMA_VERSION
    assert payload["pages"][0]["tables"][0]["rows"][1][1]["function"] == "review"
    assert sheet["B5"].value == "0013"
    assert sheet["B5"].number_format == "@"
    assert sheet["B5"].fill.fgColor.rgb.endswith("FFF2CC")


def test_pdf_digital_extrai_geometria_sem_ia():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4)
    table = Table([["A", "B"], ["001", "10"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 1, colors.black)]))
    document.build([table])

    pages = extract_pdf_tables(output.getvalue())
    cell = pages[0]["tables"][0]["rows"][0][0]
    assert pages[0]["page_width"] > 0
    assert cell["x_position"] is not None
    assert cell["width"] > 0


def _layout_payload():
    return {
        "pages": [{
            "page_number": 1,
            "page_width": 600,
            "page_height": 800,
            "regions": [
                {"value": "FICHA DE ENSAIO", "function": "title", "alignment": "center"},
                {"value": "Executante", "function": "footer"},
            ],
            "tables": [{
                "title": "Medições",
                "headers": [{"value": "Campo"}, {"value": "Valor"}],
                "rows": [
                    [{"value": "Seção", "function": "subheader", "col_span": 2}],
                    [{"value": "Número"}, {"value": "??", "ambiguous": True}],
                ],
                "column_widths": [80, 240],
            }],
        }],
    }


def _docx_xml(payload):
    from zipfile import ZipFile

    from exportacao.services import build_docx

    return ZipFile(BytesIO(build_docx(payload).getvalue())).read("word/document.xml").decode()


def _pdf_content(payload):
    """Conteúdo das páginas do PDF: o reportlab grava em ASCII85 + Flate."""
    import base64
    import re
    import zlib

    from exportacao.services import build_pdf

    raw = build_pdf(payload).getvalue()
    parts = []
    for block in re.findall(rb"stream(.*?)endstream", raw, re.DOTALL):
        try:
            parts.append(zlib.decompress(base64.a85decode(block.strip(b"\r\n"), adobe=True)))
        except Exception:
            continue
    return b"".join(parts).decode("latin-1")


def test_word_e_pdf_usam_a_mesma_paleta_do_excel():
    from reportlab.lib.colors import HexColor
    from reportlab.lib.rl_accel import fp_str

    from exportacao.style_service import DEFAULT_PALETTE

    payload = _layout_payload()
    xml = _docx_xml(payload)
    content = _pdf_content(payload)
    for role in ("title", "header", "subheader", "review", "footer"):
        background = DEFAULT_PALETTE[role]["background"]
        assert 'w:fill="{}"'.format(background) in xml
        assert fp_str(*HexColor("#" + background).rgb()) + " rg" in content


def test_word_e_pdf_preservam_mesclagens_e_textos_fora_de_tabela():
    payload = _layout_payload()
    payload["pages"][0]["tables"][0]["headers"][0]["row_span"] = 2
    xml = _docx_xml(payload)
    content = _pdf_content(payload)

    assert 'w:gridSpan w:val="2"' in xml
    assert "w:vMerge" in xml
    assert "FICHA DE ENSAIO" in xml and "Executante" in xml
    assert "FICHA DE ENSAIO" in content and "Executante" in content

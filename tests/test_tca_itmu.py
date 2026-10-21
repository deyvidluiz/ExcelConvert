import json
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection
from openpyxl.utils.cell import range_boundaries

from documentos.forms import MultiUploadForm
from exportacao.services import build_xlsx_from_template
from modelos import tca_itmu
from modelos import tca_itmu_catalog as catalog
from modelos.models import ExcelFieldMapping

PACKAGE_DIR = Path("molds/Pacote_TCA_ITMU")
requires_package = pytest.mark.skipif(
    not (PACKAGE_DIR / catalog.WORKBOOK_FILENAME).is_file(),
    reason="o pacote de referência fica fora do Git por conter dados de um ensaio real",
)

# Dados de um ensaio fictício, no lugar dos valores reais da planilha de referência.
OLD_VALUES = {
    "B4": "A",
    "C5": "S/E ANTIGA",
    "I6": "99.0000 01",
    "D15": ">5000000",
    "C56": "Equipe anterior",
    "R56": "01/0125",
}


def _cells(sheet, reference):
    min_col, min_row, max_col, max_row = range_boundaries(reference)
    return [
        sheet.cell(row, col)
        for row in range(min_row, max_row + 1)
        for col in range(min_col, max_col + 1)
    ]


def _reference_package(tmp_path, *, old_values=None):
    """Monta um pacote equivalente ao real: rótulos travados, entradas livres e fórmulas."""
    old_values = old_values or {}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = catalog.SHEET_NAME
    sheet.protection.sheet = True
    sheet.print_area = catalog.PRINT_AREA
    unlocked = Protection(locked=False)

    for item in catalog.CONFIRMED + catalog.PENDING:
        cell = sheet[item["cell"]]
        cell.protection = unlocked
        cell.value = old_values.get(item["cell"])
    for item in catalog.PRESERVED_TEXT:
        for cell in _cells(sheet, item["range"]):
            cell.protection = unlocked
            cell.value = item["text"] if item["text"] is not None else "v1-v2"

    sheet["A4"] = "Fase: "  # rótulo impresso: travado por padrão
    sheet["D13"] = "MΩ"
    sheet["H18"] = '=IF(AND(F18<>0,G18<>0),((G18/F18))*1,"")'
    sheet["X4"] = "Temperatura"  # tabela auxiliar de correção
    sheet["Y4"] = 1.32
    sheet.merge_cells("B4:F4")

    directory = tmp_path / "pacote"
    directory.mkdir()
    workbook.save(directory / catalog.WORKBOOK_FILENAME)
    (directory / catalog.CANDIDATES_FILENAME).write_text(
        json.dumps(
            [
                {"field_id": "plan1_b4", "sheet_name": "Plan1", "cell_or_range": "B4", "example_value": "A", "approved": False},
                {"field_id": "plan1_b12", "sheet_name": "Plan1", "cell_or_range": "B12", "example_value": "UI1086I", "approved": False},
                {"field_id": "plan1_m14", "sheet_name": "Plan1", "cell_or_range": "M14", "example_value": "UR %: ", "approved": False},
                {"field_id": "plan1_o18", "sheet_name": "Plan1", "cell_or_range": "O18", "example_value": None, "approved": False},
                {"field_id": "plan1_i9", "sheet_name": "Plan1", "cell_or_range": "I9", "example_value": None, "approved": False},
            ],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return directory


@pytest.mark.django_db
def test_cadastro_mapeia_so_os_campos_confirmados_e_fica_na_selecao_de_envio(tmp_path):
    report = tca_itmu.register_tca_itmu(package_dir=_reference_package(tmp_path), activate=True)
    template = report["template"]

    assert template.fields.count() == len(catalog.CONFIRMED) + len(catalog.PENDING)
    mappings = ExcelFieldMapping.objects.filter(excel_template=report["excel_template"])
    assert mappings.count() == len(catalog.CONFIRMED)

    pending_ids = {item["field_id"] for item in catalog.PENDING}
    assert not mappings.filter(field__field_id__in=pending_ids).exists()
    for field in template.fields.filter(field_id__in=pending_ids):
        assert field.excel_cell == ""
        assert field.validation_rule.startswith("CONFIRMAR ANTES DE MAPEAR")

    mapped = template.fields.get(field_id="instalacao")
    assert (mapped.excel_sheet, mapped.excel_cell) == (catalog.SHEET_NAME, "C5")
    assert str(template) in {str(choice[1]) for choice in MultiUploadForm().fields["template"].choices}


@pytest.mark.django_db
def test_cadastro_recusa_sobrescrever_uma_versao_ja_registrada(tmp_path):
    directory = _reference_package(tmp_path)
    tca_itmu.register_tca_itmu(package_dir=directory, version="1.0.0")
    with pytest.raises(tca_itmu.PackageError, match="já está cadastrada"):
        tca_itmu.register_tca_itmu(package_dir=directory, version="1.0.0")


def test_copia_limpa_esvazia_so_os_campos_variaveis(tmp_path):
    directory = _reference_package(tmp_path, old_values=OLD_VALUES)
    workbook_bytes, cleared = tca_itmu.clean_copy(directory)
    sheet = load_workbook(directory / catalog.WORKBOOK_FILENAME)[catalog.SHEET_NAME]
    clean = load_workbook(BytesIO(workbook_bytes))[catalog.SHEET_NAME]

    assert {coordinate for _, coordinate in cleared} == set(OLD_VALUES)
    for coordinate in OLD_VALUES:
        assert sheet[coordinate].value is not None
        assert clean[coordinate].value is None
    assert clean["A4"].value == "Fase: "
    assert clean["D13"].value == "MΩ"
    assert clean["M14"].value == "UR %: "
    assert clean["E18"].value == "v1-v2"
    assert clean["X4"].value == "Temperatura"
    assert clean["H18"].value == '=IF(AND(F18<>0,G18<>0),((G18/F18))*1,"")'
    assert "B4:F4" in {str(item) for item in clean.merged_cells.ranges}
    assert clean.print_area == sheet.print_area
    assert clean.protection.sheet is True


@pytest.mark.django_db
def test_preenchimento_simulado_escreve_nos_destinos_e_preserva_formulas(tmp_path):
    report = tca_itmu.register_tca_itmu(package_dir=_reference_package(tmp_path, old_values=OLD_VALUES), activate=True)
    payload = {
        "template_fields": [
            {"field_id": "instalacao", "value": "S/E ITURAMA 1"},
            {"field_id": "numero_serie", "value": "21.9515 04"},
            {"field_id": "isolamento_cc_prim_secs", "value": ">5000000"},
            {"field_id": "executante_data", "value": "02/0626"},
        ]
    }
    sheet = load_workbook(build_xlsx_from_template(payload, report["excel_template"]), data_only=False)[catalog.SHEET_NAME]

    assert sheet["C5"].value == "S/E ITURAMA 1"
    assert sheet["I6"].value == "21.9515 04"
    assert sheet["D15"].value == ">5000000"
    assert sheet["R56"].value == "02/0626"
    assert sheet["B4"].value is None  # fase não veio na extração e não herdou o ensaio anterior
    assert sheet["H18"].value == '=IF(AND(F18<>0,G18<>0),((G18/F18))*1,"")'
    assert sheet["M14"].value == "UR %: "
    assert "B4:F4" in {str(item) for item in sheet.merged_cells.ranges}


def test_revisao_classifica_as_candidatas_do_pacote(tmp_path):
    reviewed = tca_itmu.reviewed_candidates(_reference_package(tmp_path, old_values=OLD_VALUES))
    by_cell = {item["cell_or_range"]: item for item in reviewed}

    assert by_cell["B4"]["status"] == "campo_confirmado" and by_cell["B4"]["approved"] is True
    assert by_cell["B4"]["field_id"] == "fase"
    assert by_cell["B12"]["status"] == "precisa_confirmacao" and by_cell["B12"]["approved"] is False
    assert by_cell["B12"]["note"]
    assert by_cell["M14"]["status"] == "texto_fixo"
    assert by_cell["O18"]["status"] == "bloco_em_revisao"
    assert by_cell["I9"]["status"] == "descartado"


@requires_package
def test_catalogo_confere_com_a_planilha_de_referencia_real():
    report = tca_itmu.audit(PACKAGE_DIR)

    assert report["problems"] == []
    assert report["formulas"] == 105
    assert report["merges"] == 59
    assert report["candidates"] == 437
    assert report["candidates_confirmed"] == len(catalog.CONFIRMED)
    assert report["candidates_pending"] == len(catalog.PENDING)
    assert report["fields_outside_candidates"] == []

    _, cleared = tca_itmu.clean_copy(PACKAGE_DIR)
    assert {coordinate for _, coordinate in cleared} == set(report["inputs_with_value"])


@requires_package
def test_revisao_cobre_todas_as_candidatas_reais_sem_descartar_celula_com_dado():
    reviewed = tca_itmu.reviewed_candidates(PACKAGE_DIR)

    assert len(reviewed) == 437
    assert all(item["status"] != "descartado" or item["example_value"] is None for item in reviewed)
    assert {item["field_id"] for item in reviewed if item["status"] == "campo_confirmado"} == {
        item["field_id"] for item in catalog.CONFIRMED
    }


@pytest.mark.django_db
def test_download_xlsx_usa_o_modelo_cadastrado_mesmo_com_foto_em_preto_e_branco(client, user, tmp_path):
    from django.core.files.uploadedfile import SimpleUploadedFile
    from django.urls import reverse
    from PIL import Image

    from documentos.models import Document

    report = tca_itmu.register_tca_itmu(package_dir=_reference_package(tmp_path, old_values=OLD_VALUES), activate=True)
    grayscale = BytesIO()
    Image.new("L", (900, 1200), "white").save(grayscale, format="PNG")
    document = Document.objects.create(
        user=user,
        template=report["template"],
        file=SimpleUploadedFile("ficha-pb.png", grayscale.getvalue(), content_type="image/png"),
        original_name="ficha-pb.png",
        mime_type="image/png",
        size_bytes=grayscale.tell(),
        sha256="tca-itmu-pb",
        status=Document.Status.READY,
        extraction_payload={
            "pages": [],
            "alerts": [],
            "template_fields": [{"field_id": "instalacao", "value": "S/E ITURAMA 1"}],
        },
    )
    client.force_login(user)

    response = client.get(reverse("exportacao:download", args=[document.pk, "xlsx"]))

    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content), data_only=False)[catalog.SHEET_NAME]
    assert sheet["C5"].value == "S/E ITURAMA 1"
    assert sheet["H18"].value == '=IF(AND(F18<>0,G18<>0),((G18/F18))*1,"")'
    assert sheet["A4"].value == "Fase: "
    assert "B4:F4" in {str(item) for item in sheet.merged_cells.ranges}


@pytest.mark.django_db
def test_campo_numerico_vira_numero_e_o_resto_segue_como_texto(tmp_path):
    report = tca_itmu.register_tca_itmu(package_dir=_reference_package(tmp_path), activate=True)
    payload = {
        "template_fields": [
            {"field_id": "resistencia_temperatura_ambiente", "value": "31,6"},
            {"field_id": "isolamento_ca_kv_ensaio", "value": "10"},
            {"field_id": "resistencia_1s1_1s2_serie_medida", "value": "238,7"},
            {"field_id": "isolamento_cc_prim_secs", "value": ">5000000"},
            {"field_id": "corrente_nominal_1", "value": "300x600 - 5A"},
            {"field_id": "executante_data", "value": "02/0626"},
            {"field_id": "relacao_1s1_1s2_serie_relacao_nominal", "value": "60"},
        ]
    }
    sheet = load_workbook(build_xlsx_from_template(payload, report["excel_template"]), data_only=False)[catalog.SHEET_NAME]

    # As fórmulas comparam e procuram números: texto faria K44=10 dar falso.
    assert sheet["L14"].value == 31.6
    assert sheet["K44"].value == 10
    assert sheet["K38"].value == 238.7
    # O que não é inequivocamente numérico é gravado exatamente como foi lido.
    assert sheet["D15"].value == ">5000000"
    assert sheet["I7"].value == "300x600 - 5A"
    assert sheet["R56"].value == "02/0626"
    assert sheet["I38"].value == "60"

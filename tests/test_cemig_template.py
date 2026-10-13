import io
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from documentos.models import Document
from exportacao.services import build_xlsx_from_template
from extracao.preprocessamento import crop_region, preprocess_image
from extracao.gemini_service import GeminiExtractor
from extracao.services import ExtractionService
from extracao.schemas import TemplateExtraction, TemplateFieldValue, template_payload_from_model
from modelos.models import DocumentTemplate, ExcelFieldMapping, ExcelTemplate, ReferenceValue, TemplateField
from modelos.avaliacao import evaluate_payload
from modelos.controlled_test import prepare_page_preflight
from modelos.services import register_cemig_template


def nine_page_pdf():
    out = io.BytesIO()
    pdf = canvas.Canvas(out, pagesize=A4)
    for page in range(9):
        pdf.drawString(40, 800, f"Página {page + 1}")
        pdf.showPage()
    pdf.save()
    return out.getvalue()


@pytest.mark.django_db
def test_cadastra_modelo_cemig_com_nove_paginas(tmp_path):
    reference = tmp_path / "TCs_Curvelo_2.pdf"
    reference.write_bytes(nine_page_pdf())
    template = register_cemig_template(reference_path=reference)
    assert template.page_count == 9
    assert template.reference_sha256
    assert template.fields.count() >= 30
    assert template.calibration_status == "pending"
    assert template.repeat_layout_per_page is True
    assert set(template.fields.values_list("page_number", flat=True)) == {1}


@pytest.mark.django_db
def test_preflight_da_pagina_nao_chama_gemini_e_gera_artefatos(tmp_path):
    reference = tmp_path / "TCs_Curvelo_2.pdf"
    reference.write_bytes(nine_page_pdf())
    template = register_cemig_template(reference_path=reference)
    output = tmp_path / "preflight"

    manifest, processed = prepare_page_preflight(pdf_path=reference, template=template, page_number=1, output_dir=output)

    assert processed.startswith(b"\x89PNG")
    assert manifest["real_api_call"] is False
    assert manifest["fields_total"] == 32
    assert manifest["fields_calibrated"] == 0
    assert (output / "pagina-original.png").exists()
    assert (output / "pagina-preprocessada.png").exists()
    assert (output / "pagina-regioes.png").exists()
    assert (output / "manifest.json").exists()


@pytest.mark.django_db
def test_calibracao_de_regiao_marca_modelo_parcial(tmp_path):
    reference = tmp_path / "referencia.pdf"
    reference.write_bytes(nine_page_pdf())
    template = register_cemig_template(reference_path=reference)

    call_command(
        "calibrate_template_field",
        template_code=template.code,
        template_version=template.version,
        field_id="instalacao",
        x=0.1,
        y=0.1,
        width=0.2,
        height=0.03,
        confirm_visual_review=True,
    )
    template.refresh_from_db()
    field = template.fields.get(field_id="instalacao")
    assert field.region == {"x": 0.1, "y": 0.1, "width": 0.2, "height": 0.03}
    assert template.calibration_status == "partial"


@pytest.mark.django_db
def test_resposta_do_modelo_preserva_field_id_e_valor_ambiguo():
    template = DocumentTemplate.objects.create(code="teste-ficha", name="Ficha de teste", version="1.0.0", page_count=1)
    field = TemplateField.objects.create(template=template, field_id="valor", name="Valor", section="Medição", page_number=1, expected_type="decimal")
    response = TemplateExtraction(template_code=template.code, template_version=template.version, fields=[TemplateFieldValue(field_id=field.field_id, page_number=1, value="1,20", ambiguous=True, review_reason="lápis fraco")])
    payload = template_payload_from_model(response, template)
    assert payload["template_fields"][0]["field_id"] == "valor"
    assert payload["template_fields"][0]["value"] == "1,20"
    assert payload["template_fields"][0]["ambiguous"] is True


@pytest.mark.django_db
def test_gemini_estruturado_com_modelo_e_metricas_mockadas(monkeypatch):
    template = DocumentTemplate.objects.create(code="gemini-ficha", name="Ficha", version="1.0.0", page_count=1)
    TemplateField.objects.create(template=template, field_id="valor", name="Valor", section="Medição", page_number=1)
    response = SimpleNamespace(
        parsed=TemplateExtraction(template_code=template.code, template_version=template.version, fields=[TemplateFieldValue(field_id="valor", page_number=1, value="0008")]),
        usage_metadata=SimpleNamespace(prompt_token_count=200, candidates_token_count=50, total_token_count=250),
    )
    monkeypatch.setenv("GEMINI_API_KEY", "mocked-key")
    with patch("google.genai.Client") as client_class:
        client_class.return_value.models.generate_content.return_value = response
        extractor = GeminiExtractor(model="gemini-3.1-flash-lite")
        payload, tokens = extractor.extract_template(data=b"image", mime_type="image/png", document_name="ficha.png", template=template)
    assert tokens == 250
    assert extractor.last_metrics["input_tokens"] == 200
    assert payload["template_fields"][0]["value"] == "0008"
    config = client_class.return_value.models.generate_content.call_args.kwargs["config"]
    assert config.response_mime_type == "application/json"
    assert client_class.return_value.models.generate_content.call_args.kwargs["model"] == "gemini-3.1-flash-lite"


def test_preprocessamento_preserva_original_e_produz_derivado():
    from PIL import Image, ImageDraw

    original = io.BytesIO()
    image = Image.new("RGB", (1200, 1600), "white")
    ImageDraw.Draw(image).rectangle((80, 80, 1120, 1500), outline="black", width=4)
    image.save(original, format="JPEG")
    original_bytes = original.getvalue()
    processed = preprocess_image(original_bytes, "image/jpeg")
    assert original_bytes.startswith(b"\xff\xd8")
    assert processed.data.startswith(b"\x89PNG")
    assert processed.data != original_bytes
    assert 0 <= processed.quality_score <= 1
    crop = crop_region(processed.data, {"x": 0.1, "y": 0.1, "width": 0.3, "height": 0.3})
    assert crop.startswith(b"\x89PNG")


@pytest.mark.django_db
def test_pipeline_com_modelo_salva_payload_por_documento(user):
    template = DocumentTemplate.objects.create(code="teste-pipeline", name="Ficha", version="1.0.0", page_count=1)
    TemplateField.objects.create(template=template, field_id="codigo", name="Código", section="Identificação", page_number=1)
    image = io.BytesIO()
    from PIL import Image

    Image.new("RGB", (1200, 1600), "white").save(image, format="PNG")
    document = Document.objects.create(user=user, template=template, file=SimpleUploadedFile("ficha.png", image.getvalue(), content_type="image/png"), original_name="ficha.png", mime_type="image/png", size_bytes=image.tell(), sha256="template-image")

    class FakeExtractor:
        last_metrics = {"input_tokens": 120, "output_tokens": 40, "total_tokens": 160, "api_calls": 1}

        def extract(self, **kwargs):
            assert kwargs["template"].code == "teste-pipeline"
            return {"document_name": "ficha.png", "template_fields": [{"field_id": "codigo", "name": "Código", "section": "Identificação", "page_number": 1, "value": "0007", "ambiguous": False, "review_reason": None}], "template_code": "teste-pipeline", "template_version": "1.0.0", "pages": [{"page_number": 1, "tables": [{"title": "Identificação", "headers": [{"value": "Campo"}, {"value": "Valor"}], "rows": [[{"value": "Código"}, {"value": "0007", "field_id": "codigo"}]], "merged_cells": [], "alerts": []}], "alerts": []}], "alerts": []}, 160

    ExtractionService(FakeExtractor()).process(document, paid_confirmed=True)
    document.refresh_from_db()
    assert document.extraction_payload["template_fields"][0]["value"] == "0007"
    assert document.input_token_count == 120
    assert document.output_token_count == 40
    assert document.preprocessed_file


@pytest.mark.django_db
def test_xlsx_cadastrado_preserva_estilo_e_formula():
    template = DocumentTemplate.objects.create(code="teste-xlsx", name="Ficha", version="1.0.0")
    field = TemplateField.objects.create(template=template, field_id="valor", name="Valor", section="Medição", page_number=1)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dados"
    sheet["B2"] = "=1+1"
    sheet["B3"].fill = PatternFill("solid", fgColor="FF0000")
    raw = io.BytesIO()
    workbook.save(raw)
    raw.seek(0)
    excel_template = ExcelTemplate.objects.create(document_template=template, name="Destino", workbook=SimpleUploadedFile("destino.xlsx", raw.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
    ExcelFieldMapping.objects.create(excel_template=excel_template, field=field, sheet_name="Dados", cell_or_range="B3")
    output = build_xlsx_from_template({"template_fields": [{"field_id": "valor", "value": "0012"}]}, excel_template)
    result = load_workbook(output, data_only=False)
    assert result["Dados"]["B2"].value == "=1+1"
    assert result["Dados"]["B3"].value == "0012"
    assert result["Dados"]["B3"].fill.fgColor.rgb.endswith("FF0000")


@pytest.mark.django_db
def test_avaliacao_usa_referencia_humana_por_pagina_e_campo():
    template = DocumentTemplate.objects.create(code="avaliacao", name="Ficha", version="1.0.0")
    field = TemplateField.objects.create(template=template, field_id="numero", name="Número", section="Identificação", page_number=2, expected_type="decimal")
    ReferenceValue.objects.create(template=template, field=field, source_document="referencia-manual", page_number=2, value="12,50")
    run = evaluate_payload(template=template, payload={"template_fields": [{"field_id": "numero", "page_number": 2, "value": "12,50"}]})
    assert run.fields_identified == 1
    assert run.fields_correct == 1
    assert run.numeric_errors == 0

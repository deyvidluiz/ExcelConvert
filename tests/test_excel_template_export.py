from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from exportacao.services import build_xlsx_from_template
from modelos.models import DocumentTemplate, ExcelFieldMapping, ExcelTemplate, TemplateField


def _template_with_range():
    template = DocumentTemplate.objects.create(code="xlsx-range", name="Ficha", version="1.0.0")
    field = TemplateField.objects.create(
        template=template,
        field_id="leituras",
        name="Leituras",
        section="Medição",
        page_number=1,
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Plan1"
    sheet["A1"] = "=1+1"
    sheet["B3"].fill = PatternFill("solid", fgColor="FF0000")
    sheet["B3"] = "modelo"
    sheet["B4"] = "modelo"
    sheet.merge_cells("D1:E1")
    raw = BytesIO()
    workbook.save(raw)
    raw.seek(0)
    excel_template = ExcelTemplate.objects.create(
        document_template=template,
        name="Destino",
        workbook=SimpleUploadedFile(
            "destino.xlsx",
            raw.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
    )
    ExcelFieldMapping.objects.create(
        excel_template=excel_template,
        field=field,
        sheet_name="Plan1",
        cell_or_range="B3:B4",
    )
    return excel_template


@pytest.mark.django_db
def test_intervalo_preserva_formula_estilo_mesclagem_e_dimensoes():
    excel_template = _template_with_range()
    output = build_xlsx_from_template(
        {"template_fields": [{"field_id": "leituras", "value": ["0012", "0013"]}]},
        excel_template,
    )
    result = load_workbook(output, data_only=False)
    sheet = result["Plan1"]
    assert sheet["A1"].value == "=1+1"
    assert sheet["B3"].value == "0012"
    assert sheet["B4"].value == "0013"
    assert sheet["B3"].fill.fgColor.rgb.endswith("FF0000")
    assert "D1:E1" in {str(item) for item in sheet.merged_cells.ranges}


@pytest.mark.django_db
def test_intervalo_rejeita_valor_escalar_para_nao_escrever_so_a_primeira_celula():
    excel_template = _template_with_range()
    with pytest.raises(ValueError, match="lista estruturada"):
        build_xlsx_from_template(
            {"template_fields": [{"field_id": "leituras", "value": "0012"}]},
            excel_template,
        )


@pytest.mark.django_db
def test_mapeamento_de_formula_e_rejeitado():
    excel_template = _template_with_range()
    mapping = excel_template.mappings.get()
    mapping.cell_or_range = "A1"
    mapping.save(update_fields=["cell_or_range"])
    with pytest.raises(ValueError, match="fórmula"):
        build_xlsx_from_template(
            {"template_fields": [{"field_id": "leituras", "value": "0012"}]},
            excel_template,
        )
